from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_bytes(text='', password=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        lines = ['BT /F1 11 Tf 16 TL 40 795 Td']
        for line in text.splitlines():
            safe = line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            lines.append(f'({safe}) Tj T*')
        stream = DecodedStreamObject(); stream.set_data(('\n'.join([*lines, 'ET'])).encode('latin-1'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = BytesIO(); writer.write(output)
    return output.getvalue()


def docx_bytes(name='Dana Brooks', body='Platform Engineer', table='AWS EKS and Terraform'):
    document = Document()
    document.sections[0].header.paragraphs[0].text = name
    document.add_paragraph(body)
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = 'Experience', table
    document.sections[0].footer.paragraphs[0].text = 'Resume footer'
    output = BytesIO(); document.save(output)
    return output.getvalue()
