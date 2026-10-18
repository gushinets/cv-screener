import asyncio
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
with patch('search.search.Search'):
    from server import create_app
    import cv_screener as core
from history_store import HistoryStore
from common.schema import Doc
from common.store import DocumentStoreBase
from providers.store.qdrant import Store
from resumes import MAX_BYTES, ResumeError, extract_text, resume_id, validate_name
from resume_fixtures import docx_bytes, pdf_bytes


class ExtractionTests(unittest.TestCase):
    def test_word_textboxes_in_body_and_header_are_not_silently_lost(self):
        from io import BytesIO
        from docx import Document
        from docx.oxml import parse_xml
        document = Document()
        document.add_paragraph('Mikhail Gushinets')
        for container, text in [(document, 'Memory consumption reduced by 60-77%.'),
                                (document.sections[0].header, 'Release team: 6 developers, 4 QA, 1 build engineer.')]:
            paragraph = container.add_paragraph()
            paragraph._p.append(parse_xml(
                '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox>'
                f'<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>'
                '</v:textbox></v:shape></w:pict></w:r>'))
        output = BytesIO(); document.save(output)
        extracted = extract_text(output.getvalue(), '.docx')
        self.assertIn('Memory consumption reduced by 60-77%.', extracted)
        self.assertIn('Release team: 6 developers, 4 QA, 1 build engineer.', extracted)

    def test_all_formats_include_resume_text_and_word_tables_headers_and_footers(self):
        self.assertIn('Built EKS with Terraform.', extract_text(pdf_bytes('Morgan Reed\nBuilt EKS with Terraform.'), '.pdf'))
        word = extract_text(docx_bytes(), '.docx')
        for text in ['Dana Brooks', 'Platform Engineer', 'AWS EKS and Terraform', 'Resume footer']:
            self.assertIn(text, word)
        for encoding in ['utf-8-sig', 'utf-16', 'cp1251']:
            self.assertEqual(extract_text('Ирина — Python-разработчик'.encode(encoding), '.txt'), 'Ирина — Python-разработчик')

    def test_scans_passwords_corruption_binary_and_empty_files_are_rejected(self):
        for data, extension, message in [
            (pdf_bytes(), '.pdf', 'OCR'), (pdf_bytes('Morgan Reed', 'secret'), '.pdf', 'паролем'),
            (b'broken', '.pdf', 'прочитать'), (b'broken', '.docx', 'прочитать'),
            (b'\x00\x01Not a resume', '.txt', 'двоичные'), (b'', '.txt', 'пуст'),
            (b'   ', '.txt', 'не найден'), (b'x' * (MAX_BYTES + 1), '.txt', '10 МБ'),
        ]:
            with self.subTest(extension=extension, message=message), self.assertRaisesRegex(ResumeError, message):
                extract_text(data, extension)

    def test_names_cannot_escape_the_data_directory_or_target_windows_devices(self):
        for name in ['../secret.txt', 'folder\\resume.txt', 'a:stream.txt', 'CON.txt', 'NUL.pdf', 'a.txt.', ' bad.txt', 'a.exe', 'a\n.txt']:
            with self.subTest(name=name), self.assertRaises(ResumeError):
                validate_name(name)
        self.assertEqual(validate_name('Анна (резюме).PDF'), 'Анна (резюме).PDF')


class FakeStore:
    def __init__(self):
        self.documents = {}
        self.fail = False
        self.entered = self.release = None
        self.remove_entered = self.remove_release = None

    async def agetPaths(self, limit=1000):
        return {item['parent']: key for key, item in self.documents.items()}

    async def aadd(self, **document):
        self.documents[document['objectId']] = document
        if self.entered:
            self.entered.set()
            await asyncio.to_thread(self.release.wait, 5)
        if self.fail:
            raise ConnectionError('Qdrant test failure')

    async def aremove(self, identifiers):
        if self.remove_entered:
            self.remove_entered.set()
            await asyncio.to_thread(self.remove_release.wait, 5)
        for identifier in identifiers:
            self.documents.pop(identifier, None)


class UploadApiTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.data = Path(self.folder.name) / 'data'; self.data.mkdir()
        self.history = HistoryStore(Path(self.folder.name) / 'history.sqlite3')
        self.store = FakeStore()
        self.data_patch = patch('server.DATA', self.data); self.data_patch.start()
        self.store_patch = patch.object(core.search, 'getStore', return_value=self.store); self.store_patch.start()
        self.client = TestClient(create_app(self.history), base_url='http://127.0.0.1'); self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.store_patch.stop(); self.data_patch.stop(); self.folder.cleanup()

    def upload(self, name, content):
        return self.client.post('/api/resume-imports', files={'file': (name, content, 'application/octet-stream')})

    def test_uploads_are_indexed_with_original_names_and_can_be_opened_as_sources(self):
        fixtures = [('Morgan.pdf', pdf_bytes('Morgan Reed\nBuilt EKS with Terraform.')),
                    ('Dana.docx', docx_bytes()), ('Ирина.txt', 'Ирина\nPython и OpenTelemetry'.encode())]
        for name, content in fixtures:
            response = self.upload(name, content)
            self.assertEqual(response.status_code, 202, response.text)
            self.assertEqual(response.json()['status'], 'queued')
            job = self.client.get('/api/resume-imports').json()[0]
            self.assertEqual(job['status'], 'ready', job)
            self.assertEqual((self.data / name).read_bytes(), content)
            source = self.client.get('/api/resume', params={'name': name}).json()
            self.assertEqual(source['name'], name)
            self.assertEqual(source['text'], extract_text(content, Path(name).suffix))
            self.assertEqual(self.store.documents[resume_id(self.data / name)]['text'], source['text'])
        boot = self.client.get('/api/bootstrap').json()
        self.assertEqual(boot['indexed_count'], 3)
        self.assertEqual({item['format'] for item in boot['library']}, {'PDF', 'DOCX', 'TXT'})
        self.assertTrue(all(item['indexed'] for item in boot['library']))
        self.assertEqual(self.history.list(), [])
        self.assertEqual(self.client.get('/api/resume', params={'name': '../aiconfig.json'}).status_code, 404)

    def test_rejected_uploads_do_not_overwrite_or_index_anything(self):
        original = b'Original resume'
        (self.data / 'Existing.txt').write_bytes(original)
        self.assertEqual(self.upload('Existing.txt', b'Changed').status_code, 409)
        self.assertEqual((self.data / 'Existing.txt').read_bytes(), original)
        for name in ['../escape.txt', 'CON.txt', 'resume.exe']:
            self.assertEqual(self.upload(name, b'Candidate').status_code, 422)
        self.assertEqual(self.upload('Large.txt', b'x' * (MAX_BYTES + 1)).status_code, 413)
        self.assertEqual(self.client.get('/api/resume-imports').json(), [])
        self.assertEqual(self.store.documents, {})

    def test_chunked_multipart_body_has_a_limit_even_without_content_length(self):
        boundary = 'test-boundary'
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="Large.txt"\r\n\r\n'.encode()
                + b'x' * (MAX_BYTES + 128 * 1024) + f'\r\n--{boundary}--\r\n'.encode())
        response = self.client.post('/api/resume-imports', content=iter([body[:1024], body[1024:]]),
                                    headers={'Content-Type': f'multipart/form-data; boundary={boundary}'})
        self.assertEqual(response.status_code, 413, response.text)
        self.assertEqual(self.store.documents, {})
        self.assertFalse((self.data / 'Large.txt').exists())

    def test_failed_extraction_and_indexing_have_clear_status_and_roll_back_new_data(self):
        response = self.upload('Scan.pdf', pdf_bytes())
        self.assertEqual(response.status_code, 202)
        self.assertIn('OCR', self.client.get('/api/resume-imports').json()[0]['error'])
        self.assertFalse((self.data / 'Scan.pdf').exists())
        self.store.fail = True
        self.upload('Retry.txt', b'Riley Stone\nPython engineer')
        job = self.client.get('/api/resume-imports').json()[0]
        self.assertEqual(job['status'], 'error')
        self.assertIn('Qdrant', job['error'])
        self.assertEqual(self.store.documents, {})
        self.assertFalse((self.data / 'Retry.txt').exists())
        self.store.fail = False
        self.upload('Retry.txt', b'Riley Stone\nPython engineer')
        self.assertEqual(self.client.get('/api/resume-imports').json()[0]['status'], 'ready')

    def test_status_is_readable_while_indexing_and_concurrent_upload_is_rejected(self):
        self.store.entered, self.store.release = threading.Event(), threading.Event()
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(self.upload, 'Pending.txt', b'Morgan Reed\nSRE engineer')
            try:
                self.assertTrue(self.store.entered.wait(5))
                job = self.client.get('/api/resume-imports').json()[0]
                self.assertEqual(job['status'], 'indexing')
                self.assertEqual(self.upload('Second.txt', b'Dana Brooks').status_code, 409)
            finally:
                self.store.release.set()
            self.assertEqual(future.result(timeout=5).status_code, 202)
        self.assertEqual(self.client.get('/api/resume-imports').json()[0]['status'], 'ready')

    def test_upload_rejects_foreign_origin_before_processing(self):
        response = self.client.post('/api/resume-imports', files={'file': ('Resume.txt', b'Morgan Reed')},
                                    headers={'Origin': 'https://unrelated.example'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get('/api/resume-imports').json(), [])

    def test_generation_waits_until_a_failed_index_has_been_rolled_back(self):
        self.store.fail = True
        self.store.remove_entered, self.store.remove_release = threading.Event(), threading.Event()
        with patch('server.analysis.analyze', AsyncMock(side_effect=ValueError('Test answer'))) as generate, ThreadPoolExecutor(max_workers=2) as executor:
            upload = executor.submit(self.upload, 'Fail.txt', b'Morgan Reed\nSRE engineer')
            try:
                self.assertTrue(self.store.remove_entered.wait(5))
                chat = executor.submit(self.client.post, '/api/chat', json={'message': 'Find SRE', 'request_id': '00000000-0000-4000-8000-000000000001'})
                with self.assertRaises(TimeoutError):
                    chat.result(timeout=0.1)
                generate.assert_not_awaited()
            finally:
                self.store.remove_release.set()
            self.assertEqual(upload.result(timeout=5).status_code, 202)
            self.assertEqual(chat.result(timeout=5).status_code, 422)
            generate.assert_awaited_once()
        self.assertEqual(self.store.documents, {})


class IndexingTests(unittest.IsolatedAsyncioTestCase):
    async def test_cpu_work_runs_off_the_api_loop_and_ready_waits_for_qdrant(self):
        main_thread = threading.get_ident()
        chunk = Doc(page_content='Morgan Reed', metadata={})
        workers = []
        def prepare(document):
            workers.append(threading.get_ident()); return [chunk]
        def embed(chunks):
            workers.append(threading.get_ident()); chunks[0].embedding = [0.1, 0.2]
        store = SimpleNamespace(_embedding=object(), preprocess=prepare, encodeDocuments=embed, aaddChunks=AsyncMock())
        await DocumentStoreBase.aadd(store, 'id', 'cv', 'Morgan.txt', -1, 'Morgan Reed', '')
        self.assertEqual(len(workers), 2)
        self.assertTrue(all(worker != main_thread for worker in workers))
        store.aaddChunks.assert_awaited_once_with([chunk])
        client = AsyncMock()
        qdrant = SimpleNamespace(_client=client, _collection='resumes', batch_size=500, payload_limit=32_000_000)
        chunk.metadata['objectId'] = 'id'
        await Store.aaddChunks(qdrant, [chunk])
        self.assertTrue(client.batch_update_points.await_args.kwargs['wait'])


if __name__ == '__main__':
    unittest.main()
